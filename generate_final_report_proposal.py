# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER

doc = docx.Document()

title = doc.add_heading('Project Proposal', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Sri Lanka Institute of Advanced Technological Education (SLIATE)').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Advanced Technological Institute - Kurunegala').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Higher National Diploma in Information Technology').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Batch-2324(FT)').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Indoor Sports Courts Booking Web Based System').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('IT4052 | ICT Project (Individual)').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Supervisor: Ms. P.G.R.N.J Gamlath').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('M.R.Reema | KUR/IT/2324/F/0113').alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()

doc.add_heading('Table of Content', level=1)

toc_items = [
    ('1. Problem Statement', '3'),
    ('2. Objectives and Goals', '4'),
    ('3. Proposed Solution / Approach', '5'),
    ('   3.1. Overview of the Proposed Solution', '5'),
    ('   3.2. Integrated Booking Mechanism', '5'),
    ('   3.3. Real-Time Availability & Conflict Prevention', '5'),
    ('   3.4. Secure Payments & Wallet System', '5'),
    ('   3.5. Admin Control & Reporting Panel', '6'),
    ('   3.6. Automated Email Notifications', '6'),
    ('4. Scope and Limitations', '7'),
    ('   4.1. Scope of the Project', '7'),
    ('   4.2. Limitations of the Project', '7'),
    ('5. Methodology', '8'),
    ('   5.1. Software Development Approach', '8'),
    ('   5.2. System Design', '8'),
    ('   5.3. Data Collection and Processing', '8'),
    ('   5.4. Availability Checking Implementation', '8'),
    ('   5.5. Booking and Cancellation Flow', '9'),
    ('   5.6. System Implementation', '9'),
    ('   5.7. Testing and Evaluation', '9'),
    ('6. Timeline and Milestones', '10'),
    ('7. Resources and Budget', '11'),
    ('   7.1. Resources', '11'),
    ('   7.1.1. Hardware Requirements', '11'),
    ('   7.1.2. Software Requirements', '11'),
    ('   7.2. Budget', '11'),
    ('8. Literature Review', '12'),
    ('   8.1. Existing Systems and Research', '12'),
    ('   8.2. Real-Time Database Technologies in Web Applications', '12'),
    ('   8.3. User Experience in Booking Systems', '12'),
    ('   8.4. Role-Based Access Control in Web Systems', '12'),
    ('   8.5. Research Gap', '13'),
    ('9. References', '14')
]

for item, page in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.25), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p.add_run(item + '\t' + page)

doc.add_page_break()

doc.add_heading('1. Problem Statement', level=1)
doc.add_paragraph('Sports complexes deal with overlapping bookings, coach scheduling conflicts, and resource management on a daily basis. Without a centralized system, administrators manage this manually through phone calls and ledgers, leading to errors. Customers do not know whether a court is available until they call, coach leave schedules are not visible, and equipment inventory is poorly tracked.')
doc.add_paragraph('Existing apps sometimes solve pieces of the puzzle, but few bring together court bookings, coach availability, equipment inventory, membership packages, wallet-based payments, and secure online payment processing into a single, unified workflow. The lack of automated email notifications means customers often do not know whether their booking was confirmed, cancelled, or rejected until they call the facility directly.')

doc.add_heading('2. Objectives and Goals', level=1)
doc.add_paragraph('Primary Objective:', style='List Bullet')
doc.add_paragraph('To provide a robust web application that automates sports complex management and empowers players to self-serve their bookings for courts, coaches, and equipment in a unified platform.')
doc.add_paragraph('Specific Objectives:', style='List Bullet')
doc.add_paragraph('Provide a real-time booking engine for courts and coaches with conflict prevention.', style='List Bullet')
doc.add_paragraph('Integrate secure payment processing using Stripe, with support for advance payments and wallet balances.', style='List Bullet')
doc.add_paragraph('Manage equipment inventory and membership packages with full CRUD admin operations.', style='List Bullet')
doc.add_paragraph('Offer an administrative dashboard with comprehensive reporting, including monthly revenue breakdowns.', style='List Bullet')
doc.add_paragraph('Automate the booking lifecycle with HTML email notifications for confirmations and cancellations via Nodemailer.', style='List Bullet')

doc.add_heading('3. Proposed Solution / Approach', level=1)
doc.add_heading('3.1. Overview of the Proposed Solution', level=2)
doc.add_paragraph('Indoorsport is a React and Node.js web application giving sports complex operators a single platform to manage their facility. It supports two roles - User and Admin - each routed to its own interface. Users can browse available courts, view coach profiles, rent equipment, and book sessions. Admins have a dashboard for operations, reporting, and management.')
doc.add_heading('3.2. Integrated Booking Mechanism', level=2)
doc.add_paragraph('The system supports time-based and package-based bookings, allowing users to select courts, choose an available time slot, and optionally add coaches and rental equipment in a streamlined flow.')
doc.add_heading('3.3. Real-Time Availability & Conflict Prevention', level=2)
doc.add_paragraph('A real-time availability engine, powered by Firebase Firestore, tracks court schedules and coach leave dates to prevent scheduling conflicts before they happen.')
doc.add_heading('3.4. Secure Payments & Wallet System', level=2)
doc.add_paragraph('Every booking flows through a Stripe-powered checkout for secure advance payments. A digital wallet system is implemented to allow refunds and credit-based payments.')
doc.add_heading('3.5. Admin Control & Reporting Panel', level=2)
doc.add_paragraph('The admin dashboard covers court/coach operations, equipment management, package management, and booking workflows. It features a comprehensive analytics module with monthly revenue breakdowns and equipment rental tracking.')
doc.add_heading('3.6. Automated Email Notifications', level=2)
doc.add_paragraph('Automated, professionally styled HTML emails are dispatched via Nodemailer (Gmail SMTP) at every stage of the booking lifecycle (confirmation, cancellation, rejection).')

doc.add_heading('4. Scope and Limitations', level=1)
doc.add_heading('4.1. Scope of the Project', level=2)
doc.add_paragraph('The system includes user authentication (JWT/bcrypt), court/coach/equipment browsing, real-time availability checking, Stripe checkout integration, wallet balance deduction, admin CRUD tools, interactive reporting charts, and email notifications.')
doc.add_heading('4.2. Limitations of the Project', level=2)
doc.add_paragraph('The system operates primarily as a web-based application (though responsive via Tailwind CSS), and currently focuses on single-facility management without multi-branch networking.')

doc.add_heading('5. Methodology', level=1)
doc.add_heading('5.1. Software Development Approach', level=2)
doc.add_paragraph('Development follows a component-driven strategy, enabling modular features to be designed, tested, and integrated sequentially using Agile principles.')
doc.add_heading('5.2. System Design', level=2)
doc.add_paragraph('A layered architecture is utilized: React 19 (via Vite) for the frontend, Node.js/Express.js for the REST API backend, and Cloud Firestore for scalable, real-time data storage.')
doc.add_heading('5.3. Data Collection and Processing', level=2)
doc.add_paragraph('The system collects user bookings, coach utilizations, and equipment rentals. Data is processed to generate insights such as revenue trends and booking volume analytics.')
doc.add_heading('5.4. Availability Checking Implementation', level=2)
doc.add_paragraph('The Firestore data model includes dedicated Availability and Coach Availability collections to track schedules and leave dates, preventing overlaps during API validation.')
doc.add_heading('5.5. Booking and Cancellation Flow', level=2)
doc.add_paragraph('Users pick a court, time, add-ons, and proceed to Stripe checkout. Cancellations automatically trigger a wallet refund process (if applicable) and notify the user via email.')
doc.add_heading('5.6. System Implementation', level=2)
doc.add_paragraph('Key technologies include React, Node.js, Express, Firebase Firestore, Stripe API, Nodemailer, Tailwind CSS, Multer for uploads, and bcryptjs/JWT for security.')
doc.add_heading('5.7. Testing and Evaluation', level=2)
doc.add_paragraph('Extensive testing ensures correct calculation of Stripe payments, accurate wallet deductions, conflict-free booking logic, and responsive UI performance across devices.')

doc.add_heading('6. Timeline and Milestones', level=1)
doc.add_paragraph('[Timeline and milestones will be planned according to the academic schedule and project deadlines]')

doc.add_heading('7. Resources and Budget', level=1)
doc.add_heading('7.1. Resources', level=2)
doc.add_heading('7.1.1. Hardware Requirements', level=3)
doc.add_paragraph('Development machine with minimum 8 GB RAM. End-users require any modern web browser on a smartphone, tablet, or desktop.')
doc.add_heading('7.1.2. Software Requirements', level=3)
doc.add_paragraph('React 19, Tailwind CSS 4, Node.js, Express.js 4, Firebase Admin SDK 12, Stripe API, Nodemailer 6, VS Code, Git.')
doc.add_heading('7.2. Budget', level=2)
doc.add_paragraph('The system utilizes open-source frameworks (React, Node) and generous free tiers (Firebase, Gmail SMTP), minimizing costs. The primary cost involves standard Stripe transaction fees applied during active online payments.')

doc.add_heading('8. Literature Review', level=1)
doc.add_heading('8.1. Existing Systems and Research', level=2)
doc.add_paragraph('Existing systems handle court reservations but often ignore coaching, or manage payments but lack inventory tracking. A unified approach solves this fragmentation.')
doc.add_heading('8.2. Real-Time Database Technologies in Web Applications', level=2)
doc.add_paragraph('Cloud Firestore is documented as an ideal NoSQL structure for applications requiring rapid read/write operations and real-time state synchronization.')
doc.add_heading('8.3. User Experience in Booking Systems', level=2)
doc.add_paragraph('Implementing component-driven design (React) with utility-first styling (Tailwind CSS) ensures a highly responsive and intuitive booking flow.')
doc.add_heading('8.4. Role-Based Access Control in Web Systems', level=2)
doc.add_paragraph('Using stateless session management (JWT) and secure password hashing (bcrypt) effectively separates Customer and Administrator functionalities and ensures robust security.')
doc.add_heading('8.5. Research Gap', level=2)
doc.add_paragraph('None of the existing local solutions successfully bring together court scheduling, coach availability, equipment rental, membership packages, and secure Stripe online payments in one centralized place. This system closes that gap.')

doc.add_heading('9. References', level=1)
doc.add_paragraph('1. React Documentation. (2024). React - The library for web and native user interfaces.\n2. Node.js & Express Documentation. (2024).\n3. Google LLC. (2024). Cloud Firestore documentation.\n4. Stripe API Reference. (2024). Stripe API Documentation.\n5. Nodemailer Documentation. (2024).\n6. Tailwind CSS Documentation. (2024).')

doc.save(r'C:\Users\Administrator\Desktop\Final_Proposal_From_Report.docx')
doc.save(r'C:\finalproject\Final_Proposal_From_Report.docx')

import sys
import subprocess

try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx'])
    import docx

doc = docx.Document()

doc.add_heading('Project Proposal: Indoor Sports Court Booking System', 0)

doc.add_paragraph('Project Title: Indoor Sports Court Booking Web-Based System').bold = True
doc.add_paragraph('Prepared By: [Your Name / Student ID]').bold = True
doc.add_paragraph('Supervisor: [Supervisor Name]').bold = True
doc.add_paragraph('Institution / Batch: [Institution Name] / [Batch Details]').bold = True

doc.add_heading('1. Problem Statement', level=1)
doc.add_paragraph('The current recreational sports and indoor court facilities often rely on inefficient and manual booking mechanisms. Most venues depend on walk-in customers or informal phone-based reservations. This leads to frequent scheduling conflicts, overbooking, and a generally poor customer experience.')
doc.add_paragraph('Players and sports enthusiasts face significant difficulties in checking the real-time availability of courts (e.g., Badminton, Cricket, Table Tennis). Without a centralized digital system, customers often travel to a venue only to find it fully occupied, resulting in wasted time and frustration.')
doc.add_paragraph('On the operational side, venue administrators face challenges with manual record-keeping, managing multiple bookings across various sports, and tracking revenue or analyzing booking trends. Furthermore, the absence of a structured cancellation and refund process creates disputes between players and venue management. There is a clear and pressing need for a digital, real-time indoor sports court booking platform that provides convenience to customers while improving operational efficiency for administrators.')

doc.add_heading('2. Objectives and Goals', level=1)
doc.add_heading('Primary Objective:', level=2)
doc.add_paragraph('To develop a responsive, web-based Indoor Sports Court Booking System using React and Firebase that enables customers to book, manage, and cancel game slots in real time, while providing administrators with comprehensive management capabilities.')

doc.add_heading('Specific Objectives:', level=2)
doc.add_paragraph('• User-Friendly Interface: Design a platform allowing customers to seamlessly browse available courts, check real-time availability, and complete bookings online.')
doc.add_paragraph('• Dual Booking Mechanism: Implement support for both Time-Based Booking (custom start and end times) and Package-Based Booking (predefined 2-hour, 3-hour, and 4-hour packages).')
doc.add_paragraph('• Real-Time Synchronization: Develop an availability checking system that prevents double bookings and reflects slot status immediately across all active sessions.')
doc.add_paragraph('• Cancellation Management: Allow customers to cancel bookings prior to the scheduled start time seamlessly.')
doc.add_paragraph('• Admin Control Panel: Build a secure dashboard for administrators to manage courts, packages, availability slots, and track bookings.')
doc.add_paragraph('• Automated Notifications: Provide an automated email confirmation system to notify users as soon as a booking is confirmed by the admin.')
doc.add_paragraph('• Process Transparency: Ensure absolute transparency by clearly displaying accurate pricing, booking details, and slot availability at all times.')

doc.add_heading('3. Proposed Solution / Approach', level=1)
doc.add_heading('3.1. Overview', level=2)
doc.add_paragraph('The proposed system is a web-based platform built with React for the frontend and Firebase for the backend/database. The system will allow customers to select a sport, choose a booking type, verify slot availability, and confirm their reservation through a streamlined multi-step flow. The platform will support two distinct user roles: Customers and Administrators.')

doc.add_heading('3.2. Dual Booking Mechanism', level=2)
doc.add_paragraph('• Time-Based Booking: Customers select a specific start and end time. The system dynamically calculates the duration and total price based on hourly rates.')
doc.add_paragraph('• Package-Based Booking: Customers select from predefined packages. The end time is automatically calculated based on the chosen package duration.')

doc.add_heading('3.3. Real-Time Availability Module', level=2)
doc.add_paragraph('Utilizing Firebase Firestore, slot data will be managed in real time. When a customer submits a booking request, the system instantly queries Firestore to detect overlapping confirmed bookings, preventing conflicts and displaying immediate availability results.')

doc.add_heading('3.4. Cancellation Management System', level=2)
doc.add_paragraph('Customers can cancel upcoming bookings from their personal dashboard, provided the scheduled start time has not passed. Upon cancellation, the slot is immediately released in Firestore. Administrators can also override and cancel any booking if required.')

doc.add_heading('3.5. Admin Control Panel', level=2)
doc.add_paragraph('The dedicated admin interface will manage:')
doc.add_paragraph('• Courts Management: Add, edit, or toggle the availability of specific sports courts.')
doc.add_paragraph('• Packages Management: Define and update pricing structures.')
doc.add_paragraph('• Availability Management: Set hourly slot availability and manually block/open slots for maintenance.')
doc.add_paragraph('• Bookings Management: Monitor, filter, and manage all user reservations.')

doc.add_heading('3.6. Automated Email Notifications', level=2)
doc.add_paragraph('Integrated via EmailJS (or Firebase Cloud Functions), the system will automatically dispatch confirmation emails detailing the Booking ID, court type, date, time, and total amount.')

doc.add_heading('4. Scope and Limitations', level=1)
doc.add_heading('4.1. Scope of the Project', level=2)
doc.add_paragraph('• User registration and authentication via Firebase Auth.')
doc.add_paragraph('• Real-time court browsing and availability checks.')
doc.add_paragraph('• Dual booking flows (Time-Based & Package-Based).')
doc.add_paragraph('• Multi-step reservation process.')
doc.add_paragraph('• Customer dashboard for booking management.')
doc.add_paragraph('• Comprehensive Admin CRUD capabilities for courts, packages, and bookings.')
doc.add_paragraph('• Basic visual charts for revenue and booking analysis.')

doc.add_heading('4.2. Limitations of the Project', level=2)
doc.add_paragraph('• The system is purely web-based (no dedicated mobile app, though it will be responsive).')
doc.add_paragraph('• Online payment gateway integration is excluded; payments are handled locally at the venue.')
doc.add_paragraph('• Bookings operate on hourly slots (sub-hourly granularity is not supported).')
doc.add_paragraph('• Multi-venue or multi-branch management is not supported in this version.')

doc.add_heading('5. Methodology', level=1)
doc.add_heading('5.1. Software Development Approach', level=2)
doc.add_paragraph('The project will follow an Agile Software Development Life Cycle (SDLC). Development will be divided into iterative sprints, allowing for continuous refinement, regular testing, and immediate incorporation of feedback.')

doc.add_heading('5.2. System Design', level=2)
doc.add_paragraph('• Frontend Layer: React.js to provide a dynamic and responsive user interface.')
doc.add_paragraph('• Backend & Database Layer: Firebase Firestore (NoSQL) and Firebase Authentication to handle real-time business logic and secure user management.')

doc.add_heading('5.3. Availability Checking Implementation', level=2)
doc.add_paragraph('Real-time availability relies on querying the Firestore bookings collection. If a slot matches the selected date/time and has a "Confirmed" status, it is locked. Admin-blocked slots are also excluded from the public booking pool.')

doc.add_heading('6. Timeline and Milestones', level=1)
doc.add_paragraph('• Phase 1: Requirement Analysis & UI/UX Design (2 Weeks)')
doc.add_paragraph('• Phase 2: Database Setup & Authentication (1 Week)')
doc.add_paragraph('• Phase 3: Core Booking System & Availability Engine (3 Weeks)')
doc.add_paragraph('• Phase 4: Admin Dashboard & Management Tools (2 Weeks)')
doc.add_paragraph('• Phase 5: Testing, Bug Fixing, and Refinement (2 Weeks)')
doc.add_paragraph('• Phase 6: Final Deployment & Documentation (1 Week)')

doc.add_heading('7. Resources and Budget', level=1)
doc.add_heading('7.1. Software Requirements', level=2)
doc.add_paragraph('• Frontend Framework: React.js')
doc.add_paragraph('• Backend & Database: Firebase (Firestore, Auth)')
doc.add_paragraph('• Development Tool: Visual Studio Code')
doc.add_paragraph('• Version Control: Git & GitHub')

doc.add_heading('7.2. Estimated Budget', level=2)
doc.add_paragraph('• Development Tools: VS Code, React (Free)')
doc.add_paragraph('• Firebase Platform: Spark Plan (Free tier)')
doc.add_paragraph('• Third-party Libraries: Bootstrap, Chart.js, EmailJS (Free)')
doc.add_paragraph('• Documentation & Misc: Printing, Binding, Research (~2,000 LKR)')
doc.add_paragraph('Total: ~2,000 LKR')

doc.add_heading('8. Literature Review', level=1)
doc.add_heading('8.1. Existing Systems', level=2)
doc.add_paragraph('Digital booking platforms have proven to reduce administrative overhead and increase customer satisfaction. However, a significant gap remains in the adoption of these technologies among small-to-medium indoor sports courts in local regions, which often rely on paper-based methods due to the high cost of enterprise-level solutions.')

doc.add_heading('8.2. Technologies & User Experience', level=2)
doc.add_paragraph('Studies emphasize the importance of simplified, multi-step booking flows to reduce abandonment rates. Furthermore, the combination of React.js and Firebase Firestore is widely recognized as a highly scalable stack for applications requiring high-frequency read/write operations (like real-time availability checks).')

doc.add_heading('8.3. Research Gap', level=2)
doc.add_paragraph('Existing general-purpose booking platforms fail to address the specific nuances of multi-sport indoor courts (e.g., dual time/package pricing, localized admin blocking, and integrated real-time conflict prevention) within a cost-effective, easily deployable web solution. This project explicitly targets and resolves these gaps.')

doc.save(r'C:\Users\Administrator\Desktop\Project_Proposal.docx')

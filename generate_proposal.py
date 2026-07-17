import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Add a title
title = doc.add_heading('Project Proposal', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Sri Lanka Institute of Advanced Technological Education (SLIATE)').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Advanced Technological Institute-Kurunegala').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Higher National Diploma in Information Technology').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Batch-2324(FT)').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Indoor Sports Court Booking System').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('IT4052 | ICT Project (Individual)').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Supervisor: Ms.K.G.D.De.A Wijesinghe').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('M.R. Reema | KUR/IT/2324/F/0113').alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()

doc.add_heading('Table of Content', level=1)
toc_items = [
    '1. Problem Statement',
    '2. Objectives and Goals',
    '3. Proposed Solution / Approach',
    '4. Scope and Limitations',
    '5. Methodology',
    '6. Timeline and Milestones',
    '7. Resources and Budget',
    '8. Literature Review',
    '9. References'
]
for item in toc_items:
    doc.add_paragraph(item)

doc.add_page_break()

doc.add_heading('1. Problem Statement', level=1)
doc.add_paragraph('The current recreational sports and indoor courts facilities in Sri Lanka lack an efficient and transparent booking mechanism. Most venues rely on walk-in customers or informal phone-based reservations, which result in scheduling conflicts, overbooking, and poor customer experience.')
doc.add_paragraph('Players and sports enthusiasts face difficulties in checking real-time availability of sports courts and slots for sports courts such as Cricket court, Badminton court, Table Tennis court. Without a centralized system, customers often travel to a venue only to find slots already occupied, resulting in wasted time and frustration.')
doc.add_paragraph('Venue administrators also face operational challenges including manual record-keeping, difficulty in managing multiple bookings across different game types, and an inability to track revenue or analyse booking trends over time.')
doc.add_paragraph('Furthermore, the absence of a structured cancellation and refund process creates disputes between players and venue management. There is a clear need for a digital, real-time indoor sports court booking platform that provides convenience to customers while improving operational efficiency for venue administrators.')

doc.add_heading('2. Objectives and Goals', level=1)
doc.add_paragraph('Primary Objective:', style='List Bullet')
doc.add_paragraph('To develop a web-based Indoor Sports Court Booking System using React and Firebase that enables customers to book, manage, and cancel game slots in real time, while providing administrators with full management capabilities.')
doc.add_paragraph('Specific Objectives:', style='List Bullet')
doc.add_paragraph('To design and develop a user-friendly platform that allows customers to browse available courts, check real-time slot availability, and complete bookings online.', style='List Bullet')
doc.add_paragraph('To implement a dual booking mechanism supporting both Time-Based Booking and Package-Based Booking.', style='List Bullet')
doc.add_paragraph('To develop a real-time availability checking system that prevents double bookings.', style='List Bullet')
doc.add_paragraph('To implement a cancellation management system.', style='List Bullet')
doc.add_paragraph('To build a comprehensive admin panel for venue management.', style='List Bullet')
doc.add_paragraph('To provide automated email notifications for bookings.', style='List Bullet')
doc.add_paragraph('To ensure transparency in the booking process by displaying accurate pricing and availability.', style='List Bullet')

doc.add_heading('3. Proposed Solution / Approach', level=1)
doc.add_heading('3.1. Overview of the Proposed Solution', level=2)
doc.add_paragraph('This project proposes the development of an Indoor Sports Court Booking System as a web-based platform built using React (frontend) and Firebase (backend and database). The platform supports two distinct user roles: Customers (general users) and Administrators (venue managers).')
doc.add_heading('3.2. Dual Booking Mechanism', level=2)
doc.add_paragraph('Time-Based Booking: Customers select a specific start and end time. Package-Based Booking: Customers select predefined packages.')
doc.add_heading('3.3. Real-Time Availability Module', level=2)
doc.add_paragraph('Firebase Firestore is used to manage slot data in real time, immediately reflecting availability status to prevent overlapping bookings.')
doc.add_heading('3.4. Cancellation Management System', level=2)
doc.add_paragraph('Customers can cancel bookings prior to start time, automatically releasing the slot.')
doc.add_heading('3.5. Admin Control Panel', level=2)
doc.add_paragraph('Allows management of Courts, Packages, Availability, and Bookings.')
doc.add_heading('3.6. Automated Email Notifications', level=2)
doc.add_paragraph('Email notifications for successful bookings with full details.')

doc.add_heading('4. Scope and Limitations', level=1)
doc.add_heading('4.1. Scope of the Project', level=2)
doc.add_paragraph('Customer registration, login via Firebase Authentication', style='List Bullet')
doc.add_paragraph('Courts browsing with real-time availability checking', style='List Bullet')
doc.add_paragraph('Dual booking flow (Time-Based and Package-Based)', style='List Bullet')
doc.add_paragraph('My Bookings dashboard and Automated email confirmation', style='List Bullet')
doc.add_paragraph('Admin panel with full CRUD capabilities', style='List Bullet')
doc.add_paragraph('Responsive web interface using React', style='List Bullet')
doc.add_heading('4.2. Limitations of the Project', level=2)
doc.add_paragraph('Limited to web-based access; no dedicated mobile application.', style='List Bullet')
doc.add_paragraph('Online payment gateway integration is not included.', style='List Bullet')
doc.add_paragraph('Availability operates on an hourly slot basis.', style='List Bullet')
doc.add_paragraph('Does not support multi-venue or multi-branch management.', style='List Bullet')

doc.add_heading('5. Methodology', level=1)
doc.add_heading('5.1. Software Development Approach', level=2)
doc.add_paragraph('Adopts the Agile Software Development Life Cycle (SDLC) methodology for iterative development and continuous refinement.')
doc.add_heading('5.2. System Design', level=2)
doc.add_paragraph('Frontend Layer (React), Backend Layer (Node.js), and Database Layer (Firebase Firestore).')
doc.add_heading('5.3. System Implementation', level=2)
doc.add_paragraph('Frontend: React\nBackend & Database: Node.js, Firebase\nTools: VS Code, Google Chrome, Git & GitHub')

doc.add_heading('6. Timeline and Milestones', level=1)
doc.add_paragraph('[Timeline and milestones to be added based on the academic schedule]')

doc.add_heading('7. Resources and Budget', level=1)
doc.add_heading('7.1. Resources', level=2)
doc.add_paragraph('Hardware: Personal Computer/Laptop, Smartphone, Internet Connection.\nSoftware: React, Node.js, Firebase, VS Code, Google Chrome, Git/GitHub.')
doc.add_heading('7.2. Budget', level=2)
doc.add_paragraph('The total estimated cost is approximately 2,000 LKR (primarily for miscellaneous expenses like printing and documentation) as most software and platforms used have free tiers or are open-source.')

doc.add_heading('8. Literature Review', level=1)
doc.add_paragraph('Existing systems worldwide improve operational efficiency, but in Sri Lanka, most are paper-based or informal. Using real-time databases like Firebase and structured UX approaches (multi-step flows) fills this gap for small-to-medium indoor sports courts.')

doc.add_heading('9. References', level=1)
doc.add_paragraph('1. Firebase Documentation (2025). Cloud Firestore.\n2. React Documentation (2025).\n3. Bootstrap Team (2025). Bootstrap Documentation.\n4. EmailJS (2025). EmailJS Documentation.')

doc.save(r'C:\finalproject\Indoor_Sports_Court_Booking_Proposal.docx')
